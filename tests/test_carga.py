import sqlite3

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from src import carga as modulo_carga
from src.carga import Carga


@pytest.fixture
def tablas():
    listings = pd.DataFrame({
        "id": [1, 2, 3],
        "price": [100.0, np.nan, 300.0],
        "categoria_precio": pd.Categorical(["Medio", "Sin precio", "Alto"]),
        "has_availability": [True, False, True],
        "host_is_superhost": pd.array([True, None, False], dtype="boolean"),
        "bedrooms": pd.array([1, None, 2], dtype="Int64"),
        "listing_url": ["https://www.airbnb.com/rooms/1"] * 3,
    })
    reviews = pd.DataFrame({"listing_id": [1, 1], "id": [10, 11], "date": ["2024-01-01", "2024-02-01"],
                            "comments": ["bueno", "x" * 40_000]})
    calendar = pd.DataFrame({"listing_id": [1] * 5, "date": [f"2026-07-0{i}" for i in range(1, 6)]})
    return {"listings": listings, "reviews": reviews, "calendar": calendar}


def test_cargar_sqlite_inserta_todas_las_tablas(tmp_path, tablas):
    c = Carga(ruta_sqlite=tmp_path / "x.db", dir_xlsx=tmp_path)
    c.cargar_sqlite(tablas)
    con = sqlite3.connect(tmp_path / "x.db")
    assert con.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 3
    assert con.execute("SELECT COUNT(*) FROM calendar").fetchone()[0] == 5
    assert con.execute("SELECT categoria_precio FROM listings WHERE id=2").fetchone()[0] == "Sin precio"
    indices = {r[0] for r in con.execute("SELECT name FROM sqlite_master WHERE type='index'")}
    assert "idx_listings_id" in indices and "idx_calendar_listing_id" in indices
    con.close()


def test_cargar_sqlite_es_idempotente(tmp_path, tablas):
    c = Carga(ruta_sqlite=tmp_path / "x.db", dir_xlsx=tmp_path)
    c.cargar_sqlite(tablas)
    c.cargar_sqlite(tablas)
    con = sqlite3.connect(tmp_path / "x.db")
    assert con.execute("SELECT COUNT(*) FROM listings").fetchone()[0] == 3
    con.close()


def test_exportar_xlsx_crea_archivos_y_omite_tablas_gigantes(tmp_path, tablas, monkeypatch):
    monkeypatch.setattr(modulo_carga, "LIMITE_FILAS_EXCEL", 4)
    monkeypatch.setattr(modulo_carga, "ARCHIVOS_XLSX",
                        {"a.xlsx": ["listings", "reviews"], "b.xlsx": ["calendar"]})
    c = Carga(ruta_sqlite=tmp_path / "x.db", dir_xlsx=tmp_path)
    exportado = c.exportar_xlsx(tablas)
    wb = load_workbook(tmp_path / "a.xlsx", read_only=True)
    assert wb.sheetnames == ["listings", "reviews"]
    assert wb["listings"].max_row == 4  # encabezado + 3 filas
    wb.close()
    assert "calendar" not in exportado  # 5 filas > límite 4: se omite
    assert not (tmp_path / "b.xlsx").exists()


def test_verificar_carga_ok_y_detecta_diferencias(tmp_path, tablas, monkeypatch):
    monkeypatch.setattr(modulo_carga, "ARCHIVOS_XLSX", {"a.xlsx": ["listings", "reviews", "calendar"]})
    c = Carga(ruta_sqlite=tmp_path / "x.db", dir_xlsx=tmp_path)
    c.cargar_sqlite(tablas)
    c.exportar_xlsx(tablas)
    reporte = c.verificar_carga(tablas)
    assert all(r["ok"] for r in reporte.values())
    assert reporte["listings"] == {"esperado": 3, "sqlite": 3, "xlsx": 3, "ok": True}

    tablas["listings"] = pd.concat([tablas["listings"], tablas["listings"]])
    reporte = c.verificar_carga(tablas)
    assert reporte["listings"]["ok"] is False


def test_cargar_ejecuta_todo(tmp_path, tablas, monkeypatch):
    monkeypatch.setattr(modulo_carga, "ARCHIVOS_XLSX", {"a.xlsx": ["listings"]})
    c = Carga(ruta_sqlite=tmp_path / "x.db", dir_xlsx=tmp_path)
    reporte = c.cargar(tablas)
    assert reporte["listings"]["ok"] and reporte["calendar"]["xlsx"] is None
