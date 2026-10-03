"""Fase de CARGA del proceso ETL.

Contiene la clase :class:`Carga`, que:

1. inserta las tablas transformadas en la base SQLite ``output/airbnb_bogota.db``,
2. exporta las tablas a archivos XLSX en ``output/``,
3. verifica que la cantidad de registros cargados coincida con la esperada,
4. registra en el log cada evento del proceso.

Excel admite como máximo 1.048.576 filas por hoja. El calendario diario
(≈7 millones de filas) se carga completo en SQLite, pero en XLSX se exportan
sus resúmenes mensual y semanal; la omisión queda registrada como WARNING.
"""
import sqlite3
import time
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src import config
from src.logger import obtener_logger

# Límite de filas de datos por hoja de Excel (1.048.576 menos el encabezado)
LIMITE_FILAS_EXCEL = 1_048_575
# Longitud máxima de texto por celda en Excel
LIMITE_TEXTO_EXCEL = 32_767

# Archivo XLSX -> tablas (una hoja por tabla)
ARCHIVOS_XLSX = {
    "airbnb_listings.xlsx": ["listings", "hosts", "listing_amenities"],
    "airbnb_reviews.xlsx": ["reviews"],
    "airbnb_calendario_resumen.xlsx": ["calendar_mensual", "calendar_semanal", "calendar"],
}

# Índices de SQLite para acelerar los cruces y filtros más comunes
INDICES = {
    "listings": ["id", "host_id", "neighbourhood_cleansed"],
    "hosts": ["host_id"],
    "listing_amenities": ["listing_id", "amenity"],
    "calendar": ["listing_id", "date"],
    "calendar_mensual": ["listing_id"],
    "reviews": ["listing_id", "date"],
}


class Carga:
    """Carga las tablas limpias en SQLite y XLSX y verifica el resultado.

    Args:
        ruta_sqlite: archivo de la base SQLite de destino.
        dir_xlsx: carpeta donde se escriben los archivos XLSX.
    """

    def __init__(self, ruta_sqlite: Path = config.RUTA_SQLITE, dir_xlsx: Path = config.DIR_SALIDA):
        self.ruta_sqlite = Path(ruta_sqlite)
        self.dir_xlsx = Path(dir_xlsx)
        self.log = obtener_logger("carga")
        # tabla -> (archivo, filas) de lo exportado a XLSX en esta ejecución
        self.exportado_xlsx: dict[str, tuple[Path, int]] = {}

    # ---------------------------------------------------------------- utilidades
    @staticmethod
    def _preparar(df: pd.DataFrame) -> pd.DataFrame:
        """Convierte columnas categóricas a texto para que SQLite/Excel las acepten."""
        df = df.copy()
        for col in df.columns:
            if isinstance(df[col].dtype, pd.CategoricalDtype):
                df[col] = df[col].astype("string")
        return df

    # ------------------------------------------------------------------- SQLite
    def cargar_sqlite(self, tablas: dict[str, pd.DataFrame], chunksize: int = 100_000) -> None:
        """Inserta cada DataFrame como una tabla de SQLite (reemplaza si existe)."""
        self.ruta_sqlite.parent.mkdir(parents=True, exist_ok=True)
        self.log.info("Cargando %d tablas en SQLite: %s", len(tablas), self.ruta_sqlite)
        con = sqlite3.connect(self.ruta_sqlite)
        try:
            # Carga masiva: se desactiva el journal para acelerar la inserción
            con.execute("PRAGMA journal_mode = OFF")
            con.execute("PRAGMA synchronous = OFF")
            for nombre, df in tablas.items():
                inicio = time.perf_counter()
                self._preparar(df).to_sql(nombre, con, if_exists="replace", index=False,
                                          chunksize=chunksize)
                for col in INDICES.get(nombre, []):
                    if col in df.columns:
                        con.execute(f'CREATE INDEX IF NOT EXISTS idx_{nombre}_{col} ON "{nombre}" ("{col}")')
                con.commit()
                self.log.info("SQLite: tabla '%s' cargada con %d registros (%.1f s)",
                              nombre, len(df), time.perf_counter() - inicio)
        except Exception:
            self.log.exception("Error cargando datos en SQLite")
            raise
        finally:
            con.close()

    # --------------------------------------------------------------------- XLSX
    def exportar_xlsx(self, tablas: dict[str, pd.DataFrame]) -> dict[str, tuple[Path, int]]:
        """Exporta las tablas a XLSX según ``ARCHIVOS_XLSX`` (una hoja por tabla).

        Las tablas que superan el límite de filas de Excel se omiten con un
        WARNING; los textos más largos que el máximo de una celda se recortan.

        Returns:
            ``{tabla: (ruta_archivo, filas)}`` de las hojas escritas.
        """
        self.dir_xlsx.mkdir(parents=True, exist_ok=True)
        self.exportado_xlsx = {}
        for archivo, nombres in ARCHIVOS_XLSX.items():
            hojas = {}
            for nombre in nombres:
                if nombre not in tablas:
                    continue
                df = tablas[nombre]
                if len(df) > LIMITE_FILAS_EXCEL:
                    self.log.warning("XLSX: la tabla '%s' (%d filas) supera el límite de Excel (%d); "
                                     "se omite del XLSX y queda completa en SQLite",
                                     nombre, len(df), LIMITE_FILAS_EXCEL)
                    continue
                hojas[nombre] = self._recortar_textos(self._preparar(df), nombre)
            if not hojas:
                continue

            ruta = self.dir_xlsx / archivo
            inicio = time.perf_counter()
            try:
                with pd.ExcelWriter(ruta, engine="xlsxwriter",
                                    engine_kwargs={"options": {"strings_to_urls": False}}) as writer:
                    for nombre, df in hojas.items():
                        df.to_excel(writer, sheet_name=nombre[:31], index=False)
                        self.exportado_xlsx[nombre] = (ruta, len(df))
            except Exception:
                self.log.exception("Error exportando %s", ruta)
                raise
            self.log.info("XLSX: %s exportado con hojas %s (%.1f s)",
                          ruta.name, {n: len(d) for n, d in hojas.items()}, time.perf_counter() - inicio)
        return self.exportado_xlsx

    def _recortar_textos(self, df: pd.DataFrame, nombre: str) -> pd.DataFrame:
        """Recorta textos que superan el máximo de caracteres de una celda de Excel."""
        for col in df.columns:
            if pd.api.types.is_string_dtype(df[col]) or df[col].dtype == object:
                largos = df[col].astype("string").str.len() > LIMITE_TEXTO_EXCEL
                if largos.any():
                    self.log.warning("XLSX: %d textos de '%s.%s' recortados a %d caracteres",
                                     int(largos.sum()), nombre, col, LIMITE_TEXTO_EXCEL)
                    df.loc[largos, col] = df.loc[largos, col].astype("string").str.slice(0, LIMITE_TEXTO_EXCEL)
        return df

    # ------------------------------------------------------------- verificación
    def verificar_carga(self, tablas: dict[str, pd.DataFrame]) -> dict[str, dict]:
        """Compara los registros esperados con los cargados en SQLite y XLSX.

        Returns:
            ``{tabla: {"esperado", "sqlite", "xlsx", "ok"}}``; ``xlsx`` es
            ``None`` cuando la tabla no se exportó a Excel.
        """
        reporte = {}
        con = sqlite3.connect(self.ruta_sqlite)
        libros = {}
        try:
            for nombre, df in tablas.items():
                esperado = len(df)
                try:
                    en_sqlite = con.execute(f'SELECT COUNT(*) FROM "{nombre}"').fetchone()[0]
                except sqlite3.OperationalError:
                    en_sqlite = None
                en_xlsx = None
                if nombre in self.exportado_xlsx:
                    ruta = self.exportado_xlsx[nombre][0]
                    if ruta not in libros:
                        libros[ruta] = load_workbook(ruta, read_only=True)
                    en_xlsx = libros[ruta][nombre[:31]].max_row - 1
                ok = en_sqlite == esperado and (en_xlsx is None or en_xlsx == esperado)
                reporte[nombre] = {"esperado": esperado, "sqlite": en_sqlite, "xlsx": en_xlsx, "ok": ok}
                if ok:
                    self.log.info("Verificación OK '%s': esperado=%d, SQLite=%s, XLSX=%s",
                                  nombre, esperado, en_sqlite, en_xlsx if en_xlsx is not None else "no exportado")
                else:
                    self.log.error("Verificación FALLIDA '%s': esperado=%d, SQLite=%s, XLSX=%s",
                                   nombre, esperado, en_sqlite, en_xlsx)
        finally:
            con.close()
            for libro in libros.values():
                libro.close()
        return reporte

    # ------------------------------------------------------------------ todo
    def cargar(self, tablas: dict[str, pd.DataFrame], exportar_excel: bool = True) -> dict[str, dict]:
        """Carga en SQLite, exporta a XLSX y verifica. Devuelve el reporte de verificación."""
        self.cargar_sqlite(tablas)
        if exportar_excel:
            self.exportar_xlsx(tablas)
        else:
            self.log.warning("Exportación a XLSX omitida por parámetro")
        reporte = self.verificar_carga(tablas)
        fallidas = [n for n, r in reporte.items() if not r["ok"]]
        if fallidas:
            self.log.error("Carga con diferencias en: %s", fallidas)
        else:
            self.log.info("Carga verificada: %d tablas sin diferencias", len(reporte))
        return reporte
